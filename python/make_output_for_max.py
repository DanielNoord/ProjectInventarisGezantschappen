#!/usr/bin/env python3

import argparse
import csv
import shutil
from collections.abc import Iterator
from pathlib import Path

from contentdm.file_scan_mapper import FileIdentifier, FileScanMapper
from contentdm.models import FileRow, SeriesRow
from contentdm.sqlite_maker import SqliteDatabaseMaker
from data_parsing.load_database import initialize_database_for_xml
from openpyxl import load_workbook
from typing_utils.translations_classes import Database
from xlsx_make import create_sanitized_xlsx

_EXTERNAL_DRIVE = Path("/Volumes") / "LaCie"
_INPUTS = Path("inputs")
_OUTPUTS = _EXTERNAL_DRIVE / "outputs"
_SCANS = _EXTERNAL_DRIVE / "VolumesLegazione"


def _get_all_griglie(sanitized_dir: Path) -> Iterator[Path]:
    """Yields all file paths to the santizied griglie in sorted order."""
    files = [
        (i, i.name.replace("Paesi Bassi VOLUME", "").replace("_it_IT.xlsx", ""))
        for i in sanitized_dir.iterdir()
        if i.name.startswith("Paesi")
    ]

    for file, _ in sorted(files, key=lambda x: int(x[1])):
        yield file


def _get_all_files(griglie: Path, database: Database) -> Iterator[FileRow | SeriesRow]:
    files_and_scans = FileScanMapper().run(griglie)

    workbook = load_workbook(griglie)
    sheet = workbook[workbook.sheetnames[0]]

    for index, row in enumerate(sheet.iter_rows()):
        row_title = FileIdentifier(str(row[0].value))

        if row_title.endswith("_title"):
            yield SeriesRow.from_row(row, database)

            if index == 0:
                series_name = row_title.removesuffix("_title")
                for suffix in ["_d", "_p"]:
                    name = series_name + suffix
                    row[0].value = name
                    yield FileRow.from_row(row, [name + ".tif"], database)

        elif row_title in files_and_scans:
            yield FileRow.from_row(row, files_and_scans[row_title], database)

    print(f"Finished parsing all {index + 1} rows in {griglie.stem}")


class ContentDMFileWriter:
    """Class which can write tab-delimited txt files to be imported by ContentDM."""

    def __init__(self, *, sanitize: bool, copy_scans: bool, skip_contentdm: bool) -> None:
        # Load the database with translations and individuals
        self._database = initialize_database_for_xml()
        """Database with all translations and individuals."""

        self._sanitized_dir = Path("outputs") / "VolumesExcelSanitized" / "it_IT"
        """Directory with input .xlsx files."""

        # Sanitize the input .xlsx files
        if sanitize:
            create_sanitized_xlsx(str(_INPUTS / "VolumesExcel" / "it_IT"))
            print("Sanitized all input Excel files.")

        self._copy_scans = copy_scans
        """Whether to copy the scan files to the output directory."""

        self._skip_contentdm = skip_contentdm
        """Whether to skip creating ContentDM files."""

        if self._copy_scans and self._skip_contentdm:
            raise ValueError("Cannot copy scans when skipping ContentDM file creation.")

    def run(self) -> None:
        """Create a TSV file for ContentDM."""
        output_dir = _OUTPUTS / "contentdm"

        with SqliteDatabaseMaker(_OUTPUTS) as sqlite_maker:
            sqlite_maker.insert_individuals(self._database.individuals)

            for griglie in _get_all_griglie(self._sanitized_dir):
                with sqlite_maker.griglie() as griglie_inserter:
                    (output_dir / griglie.stem).mkdir(parents=True, exist_ok=True)

                    for file_or_serie in _get_all_files(griglie, self._database):
                        if isinstance(file_or_serie, SeriesRow):
                            griglie_inserter.add_series(file_or_serie)
                            continue

                        if not file_or_serie.title:
                            raise ValueError(f"No title found for file {file_or_serie.file_id}")

                        griglie_inserter.add_file(file_or_serie)

                        row = {
                            "filename": file_or_serie.file_id,
                            "title_en_gb": file_or_serie.title_en,
                            "scanname": "",
                        }

                        if self._skip_contentdm:
                            continue

                        # Only do the filesystem operations if necessary
                        file_dir = output_dir / griglie.stem / file_or_serie.file_id
                        file_dir.mkdir(parents=True, exist_ok=True)

                        # Only create scans directory if we need to copy scans
                        if self._copy_scans:
                            (file_dir / "scans").mkdir(parents=True, exist_ok=True)

                        with open(
                            file_dir / f"{file_or_serie.file_id}.txt", "w", encoding="utf-8"
                        ) as tsv_file:
                            tsv_writer = csv.DictWriter(
                                tsv_file,
                                dialect=csv.excel_tab,
                                fieldnames=list(row),
                            )
                            tsv_writer.writeheader()
                            tsv_writer.writerow(row)

                            for scan in file_or_serie.scans:
                                row["scanname"] = scan
                                tsv_writer.writerow(row)

                                if self._copy_scans:
                                    shutil.copy(
                                        (
                                            _SCANS
                                            / scan.split("_")[0].upper()
                                            / scan.replace("ms", "MS")
                                        ),
                                        file_dir / "scans" / scan,
                                    )


def main() -> None:
    """Main function to run the ContentDM file writer."""
    argparser = argparse.ArgumentParser(description="Bla di bla.")
    argparser.add_argument(
        "--sanitize",
        action="store_true",
        help="Sanitize the input Excel files before processing.",
    )
    argparser.add_argument(
        "--copy-scans",
        action="store_true",
        help="Copy the scan files to the output directory.",
    )
    argparser.add_argument(
        "--skip-contentdm",
        action="store_true",
        help="Skip creating ContentDM files.",
    )
    args = argparser.parse_args()
    maker = ContentDMFileWriter(
        sanitize=args.sanitize, copy_scans=args.copy_scans, skip_contentdm=args.skip_contentdm
    )
    maker.run()


if __name__ == "__main__":
    main()
